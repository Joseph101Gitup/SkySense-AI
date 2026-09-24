"""
SKYsense AI — Professional PowerPoint Presentation Generator
Master's Degree Final Project Presentation
Uses only verified, empirical data. Never fabricates results.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.dml import MSO_THEME_COLOR
import copy
from pathlib import Path
from datetime import datetime

# ── Output path ──────────────────────────────────────────────────────────────
OUTPUT_PATH = Path(__file__).resolve().parent.parent / "SKYsense_AI_Presentation.pptx"

# ── Brand Palette ─────────────────────────────────────────────────────────────
NAVY        = RGBColor(0x0A, 0x0F, 0x2E)   # slide background / dark surfaces
DEEP_BLUE   = RGBColor(0x08, 0x2A, 0x5E)   # section blocks
SKY_BLUE    = RGBColor(0x00, 0x8B, 0xD8)   # primary accent
CYAN        = RGBColor(0x00, 0xD4, 0xFF)   # highlight lines
GREEN       = RGBColor(0x00, 0xC8, 0x8E)   # success / correct
AMBER       = RGBColor(0xFF, 0xA5, 0x00)   # warning / future
ROSE        = RGBColor(0xFF, 0x3E, 0x6E)   # danger / heavy rain
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY  = RGBColor(0xB0, 0xBE, 0xD4)
DARK_CARD   = RGBColor(0x10, 0x1A, 0x3A)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]  # blank layout

# ── Helpers ───────────────────────────────────────────────────────────────────

def new_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color=NAVY):
    """Fill slide background with solid color."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill_color=DEEP_BLUE, alpha=None):
    """Add a filled rectangle."""
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.line.fill.background()  # no border
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    return shape


def txbox(slide, l, t, w, h, text, size=18, bold=False,
          color=WHITE, align=PP_ALIGN.LEFT, italic=False, wrap=True):
    """Add a text box."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def add_para(tf, text, size=14, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, italic=False, space_before=0):
    from pptx.util import Pt as _Pt
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = _Pt(space_before)
    if text:
        run = p.add_run()
        run.text = text
        run.font.size = _Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return p


def hline(slide, t, color=CYAN, l=0.4, w=12.5, thickness=0.03):
    """Horizontal accent line."""
    shape = slide.shapes.add_shape(1,
        Inches(l), Inches(t), Inches(w), Inches(thickness))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def slide_number(slide, n, total=18):
    txbox(slide, 12.0, 7.15, 1.2, 0.3,
          f"{n} / {total}", size=9, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)


def section_badge(slide, label, l=0.4, t=0.15, w=2.5, h=0.35):
    rect(slide, l, t, w, h, SKY_BLUE)
    txbox(slide, l+0.08, t+0.02, w-0.1, h-0.04,
          label, size=10, bold=True, color=WHITE)


def slide_title(slide, title, subtitle=None, t=0.55):
    txbox(slide, 0.4, t, 12.5, 0.7, title,
          size=30, bold=True, color=WHITE)
    if subtitle:
        txbox(slide, 0.4, t+0.72, 12.5, 0.4, subtitle,
              size=15, color=CYAN, italic=True)


def bullet_box(slide, items, l, t, w, h, size=14, color=WHITE, gap=0.3):
    """Each item is a (bullet_char, text) tuple."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for (bullet, text) in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(gap * 72)  # points
        run = p.add_run()
        run.text = f"{bullet}  {text}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return tb


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1 — Title Slide
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
rect(s, 0, 0, 13.33, 2.2, DEEP_BLUE)
rect(s, 0, 2.17, 13.33, 0.07, CYAN)

txbox(s, 0.5, 0.3, 12.3, 0.55, "MASTER'S DEGREE FINAL PROJECT",
      size=13, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
txbox(s, 0.5, 0.8, 12.3, 1.1,
      "SKYsense AI",
      size=58, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txbox(s, 0.5, 1.85, 12.3, 0.4,
      "Precision Cloud & Rainfall Intelligence via Deep Transfer Learning",
      size=17, color=LIGHT_GREY, align=PP_ALIGN.CENTER, italic=True)

# Three metric cards
card_data = [
    ("2,543", "Cloud Specimens"),
    ("58.27%", "Test Accuracy"),
    ("87", "Automated Tests"),
]
for i, (val, lbl) in enumerate(card_data):
    cx = 1.5 + i * 3.8
    rect(s, cx, 2.6, 3.0, 1.4, DARK_CARD)
    rect(s, cx, 2.6, 3.0, 0.06, SKY_BLUE)
    txbox(s, cx, 2.72, 3.0, 0.7, val,
          size=32, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    txbox(s, cx, 3.35, 3.0, 0.4, lbl,
          size=12, color=LIGHT_GREY, align=PP_ALIGN.CENTER)

txbox(s, 0.5, 4.25, 12.3, 0.4,
      "Technology Stack:  Python 3.11  ·  TensorFlow 2.21  ·  Keras 3  ·  Xception CNN  ·  Django 5.2  ·  Chart.js",
      size=12, color=LIGHT_GREY, align=PP_ALIGN.CENTER)

txbox(s, 0.5, 4.75, 12.3, 0.35,
      "Department of Computer Science  ·  2026",
      size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

hline(s, 7.1, CYAN, w=12.5)
txbox(s, 0.5, 7.15, 12.3, 0.3,
      "All performance metrics are empirical values measured on independent held-out test data. No results have been fabricated.",
      size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER, italic=True)
slide_number(s, 1)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2 — Problem Statement
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "PROBLEM STATEMENT")
slide_title(s, "The Challenge: Localized Rainfall Prediction",
            "Traditional methods are expensive, coarse-grained, and inaccessible")
hline(s, 1.55)

# Left column — Problem
rect(s, 0.4, 1.7, 5.8, 4.9, DARK_CARD)
rect(s, 0.4, 1.7, 5.8, 0.07, ROSE)
txbox(s, 0.5, 1.75, 5.6, 0.4, "❌  Current Limitations",
      size=14, bold=True, color=ROSE)
items = [
    ("▸", "Doppler radar networks cost $1M–$5M per unit"),
    ("▸", "Satellite imagery has 1–6 km spatial resolution gaps"),
    ("▸", "Rural and developing regions have zero radar coverage"),
    ("▸", "Forecast latency of 15–60 minutes for localized events"),
    ("▸", "No affordable edge-deployable alternative exists"),
    ("▸", "Ground-truth optical data is abundant but unused"),
]
bullet_box(s, items, 0.5, 2.2, 5.6, 4.2, size=13, gap=0.18)

# Right column — Opportunity
rect(s, 6.6, 1.7, 6.3, 4.9, DARK_CARD)
rect(s, 6.6, 1.7, 6.3, 0.07, GREEN)
txbox(s, 6.7, 1.75, 6.1, 0.4, "✅  The Opportunity",
      size=14, bold=True, color=GREEN)
items2 = [
    ("▸", "Ground-level sky cameras are low-cost (< $50)"),
    ("▸", "Cloud morphology carries strong precipitation signals"),
    ("▸", "Deep CNNs extract spatial texture features automatically"),
    ("▸", "Transfer learning overcomes small dataset constraints"),
    ("▸", "Edge TFLite models can run on Raspberry Pi in < 320 ms"),
    ("▸", "REST API allows any sensor node to submit optical data"),
]
bullet_box(s, items2, 6.7, 2.2, 6.1, 4.2, size=13, color=WHITE, gap=0.18)

slide_number(s, 2)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3 — Objective & Scope
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "OBJECTIVE & SCOPE")
slide_title(s, "Research Objective",
            "Build and validate a deep learning system for cloud-image-based rainfall classification")
hline(s, 1.55)

# Three objective columns
goals = [
    (SKY_BLUE, "01", "AI Classification Engine",
     ["Train Xception CNN on 2,543 CCSN cloud images",
      "Two-phase transfer learning methodology",
      "Classify into 3 rainfall probability regimes",
      "Export validated TFLite edge model (20.5 MB)"]),
    (GREEN, "02", "Django Web Platform",
     ["Authenticated web interface with prediction history",
      "Real-time inference with confidence visualization",
      "Projector-ready demo mode for presentations",
      "Headless REST API for automated ingestion"]),
    (AMBER, "03", "IoT-Ready Architecture",
     ["Forward-compatible REST API specification",
      "Simulated edge device client (ESP32 / RPi)",
      "Optional telemetry fields: GPS, temp, humidity",
      "Software layer complete — hardware is future work"]),
]
for i, (color, num, title, pts) in enumerate(goals):
    cx = 0.4 + i * 4.3
    rect(s, cx, 1.75, 4.1, 5.3, DARK_CARD)
    rect(s, cx, 1.75, 4.1, 0.07, color)
    txbox(s, cx+0.1, 1.8, 4.0, 0.55, num,
          size=28, bold=True, color=color)
    txbox(s, cx+0.1, 2.35, 4.0, 0.45, title,
          size=14, bold=True, color=WHITE)
    for j, pt in enumerate(pts):
        txbox(s, cx+0.1, 2.9 + j*0.5, 3.9, 0.45,
              f"▸  {pt}", size=12, color=LIGHT_GREY)

slide_number(s, 3)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 4 — Dataset
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "DATASET")
slide_title(s, "CCSN Cloud Database — 2,543 Verified Specimens",
            "Cirrus Cumulus Stratus Nimbus Database · Kaggle · WMO-compliant cloud genus classification")
hline(s, 1.55)

# Partition table
rect(s, 0.4, 1.7, 8.0, 4.95, DARK_CARD)
headers = ["Class", "Cloud Genera (WMO)", "Train (70%)", "Val (15%)", "Test (15%)", "Total"]
col_x = [0.5, 1.6, 5.2, 6.15, 7.1, 7.9]
col_w = [1.1, 3.6, 0.9, 0.9, 0.9, 1.1]

# Header row
rect(s, 0.4, 1.7, 8.0, 0.4, SKY_BLUE)
for j, h in enumerate(headers):
    txbox(s, col_x[j], 1.73, col_w[j], 0.35, h,
          size=11, bold=True, color=WHITE)

rows = [
    ("Low/Med Rain", "As, St, Sc, Ns",          "701",  "151", "150", "1,002"),
    ("Med/Heavy Rain","Cb, Cu",                  "435",  "93",  "93",  "621"),
    ("No/Low Rain",  "Ci, Cs, Cc, Ac, Ct",       "643",  "139", "138", "920"),
    ("TOTAL",        "11 Genera → 3 Classes",   "1,779","383", "381", "2,543"),
]
row_colors = [DARK_CARD, RGBColor(0x12,0x1F,0x40), DARK_CARD, DEEP_BLUE]
for i, (row, rc) in enumerate(zip(rows, row_colors)):
    ry = 2.15 + i * 0.62
    rect(s, 0.4, ry, 8.0, 0.6, rc)
    for j, val in enumerate(row):
        bold = (i == 3) or (j == 0)
        col = CYAN if i == 3 else WHITE
        txbox(s, col_x[j], ry+0.08, col_w[j], 0.44,
              val, size=12, bold=bold, color=col)

txbox(s, 0.4, 4.65, 8.0, 0.3,
      "★ Stratified 70/15/15 split · seed=42 · Zero overlap verified (SHA-256 file hashes)",
      size=10, italic=True, color=LIGHT_GREY)

# Side notes
notes = [
    (SKY_BLUE,  "Reproducible",  "seed = 42\nDeterministic\nstratified split"),
    (GREEN,     "Zero Leakage",  "Train ∩ Val ∩ Test\n= ∅  verified"),
    (AMBER,     "No Augment",    "Val & Test receive\nrescaling only"),
]
for i, (c, hdr, body) in enumerate(notes):
    ry = 1.7 + i * 1.6
    rect(s, 8.7, ry, 4.2, 1.4, DARK_CARD)
    rect(s, 8.7, ry, 4.2, 0.06, c)
    txbox(s, 8.8, ry+0.1, 4.0, 0.35, hdr, size=13, bold=True, color=c)
    txbox(s, 8.8, ry+0.48, 4.0, 0.85, body, size=12, color=WHITE)

slide_number(s, 4)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 5 — AI Model Architecture
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "AI MODEL")
slide_title(s, "Xception CNN — Deep Transfer Learning Architecture",
            "Extreme Inception · 20.8M Parameters · Depthwise Separable Convolutions")
hline(s, 1.55)

# Pipeline boxes
pipeline = [
    (SKY_BLUE, "INPUT", "256×256×3\nRGB Tensor\n[0, 1] normalized"),
    (DEEP_BLUE,"ENTRY\nFLOW",  "3× Depthwise\nSep. Conv Blocks\n+ Max Pooling"),
    (DEEP_BLUE,"MIDDLE\nFLOW", "8× Residual\nBlocks\n728 filters"),
    (DEEP_BLUE,"EXIT\nFLOW",   "1024→1536\n→2048 filters\n+ Max Pooling"),
    (SKY_BLUE, "GAP",          "Global Average\nPooling\n2048-dim vector"),
    (RGBColor(0x1A,0x3A,0x20),"DROPOUT","rate = 0.30\nRegularization"),
    (GREEN,    "OUTPUT",       "Dense(3)\nSoftmax\nProbabilities"),
]
bw, bh = 1.52, 1.8
start_x = 0.35
for i, (c, title, body) in enumerate(pipeline):
    bx = start_x + i * (bw + 0.1)
    rect(s, bx, 1.85, bw, bh, DARK_CARD)
    rect(s, bx, 1.85, bw, 0.07, c)
    txbox(s, bx+0.05, 1.9, bw-0.1, 0.45, title,
          size=11, bold=True, color=c, align=PP_ALIGN.CENTER)
    txbox(s, bx+0.05, 2.4, bw-0.1, 1.2, body,
          size=10, color=WHITE, align=PP_ALIGN.CENTER)
    if i < len(pipeline) - 1:
        txbox(s, bx+bw+0.01, 2.55, 0.12, 0.4, "▶",
              size=12, color=CYAN, align=PP_ALIGN.CENTER)

# Why Xception
rect(s, 0.35, 3.85, 12.6, 1.8, DARK_CARD)
txbox(s, 0.45, 3.9, 12.4, 0.35, "Why Xception over ResNet / VGG?",
      size=13, bold=True, color=CYAN)
reasons = [
    "▸  Depthwise Separable Convolutions decouple spatial and channel correlations — superior for cloud texture analysis",
    "▸  8× fewer multiplications than standard convolutions while maintaining representational power",
    "▸  ImageNet pre-training provides low-level edge detectors transferable to cloud morphology",
    "▸  Compact edge export: full Keras model 80.1 MB → TFLite 20.5 MB → 40ms CPU inference",
]
for i, r in enumerate(reasons):
    txbox(s, 0.45, 4.3 + i*0.33, 12.4, 0.3, r, size=11, color=WHITE)

slide_number(s, 5)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 6 — Transfer Learning Methodology
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "TRAINING METHODOLOGY")
slide_title(s, "Two-Phase Transfer Learning",
            "ImageNet Pre-training → Feature Extraction → Fine-Tuning")
hline(s, 1.55)

for i, (phase_color, phase_num, phase_name, details) in enumerate([
    (SKY_BLUE, "PHASE 1", "Feature Extraction",
     ["Backbone: All 132 Xception layers FROZEN",
      "Only 6,147 classification head parameters trained",
      "Optimizer: Adam  ·  LR = 1×10⁻³",
      "Epochs: 10  ·  Batch: 32",
      "Goal: Adapt ImageNet features to cloud domain"]),
    (GREEN, "PHASE 2", "Fine-Tuning",
     ["Top 30 Xception layers UNFROZEN (≈4.8M params)",
      "Conservative LR = 1×10⁻⁵ to prevent catastrophic forgetting",
      "EarlyStopping (patience=5)  ·  ReduceLROnPlateau (factor=0.5)",
      "Checkpoint: best val_loss → best_xception_rainfall.keras",
      "Epochs: 25  ·  Best weights auto-restored"]),
]):
    cx = 0.4 + i * 6.5
    rect(s, cx, 1.7, 6.2, 5.3, DARK_CARD)
    rect(s, cx, 1.7, 6.2, 0.08, phase_color)
    txbox(s, cx+0.15, 1.75, 6.0, 0.4, phase_num,
          size=11, bold=True, color=phase_color)
    txbox(s, cx+0.15, 2.1, 6.0, 0.5, phase_name,
          size=20, bold=True, color=WHITE)
    hline(s, 2.65, phase_color, l=cx+0.15, w=5.8, thickness=0.025)
    for j, d in enumerate(details):
        txbox(s, cx+0.2, 2.75 + j*0.55, 5.9, 0.5,
              f"▸  {d}", size=13, color=WHITE)

txbox(s, 0.4, 7.1, 12.5, 0.3,
      "★  Reproducible: RANDOM_SEED = 42  ·  Val & Test partitions never augmented  ·  Test set strictly held-out until final evaluation",
      size=10, italic=True, color=LIGHT_GREY)

slide_number(s, 6)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 7 — Empirical Results
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "EMPIRICAL RESULTS")
slide_title(s, "Model Performance — Independent Test Set (N = 381)",
            "All metrics are empirical. Evaluated once on strictly held-out data never seen during training.")
hline(s, 1.55)

# Main metric cards row
metrics = [
    (CYAN,  "58.27%", "Test Accuracy"),
    (AMBER, "0.9103", "Categorical Loss"),
    (GREEN, "56.60%", "Macro F1-Score"),
    (SKY_BLUE,"57.58%","Weighted F1"),
    (ROSE,  "381",    "Test Specimens"),
]
for i, (c, val, lbl) in enumerate(metrics):
    cx = 0.35 + i * 2.55
    rect(s, cx, 1.7, 2.4, 1.3, DARK_CARD)
    rect(s, cx, 1.7, 2.4, 0.06, c)
    txbox(s, cx, 1.8, 2.4, 0.65, val,
          size=26, bold=True, color=c, align=PP_ALIGN.CENTER)
    txbox(s, cx, 2.38, 2.4, 0.35, lbl,
          size=11, color=LIGHT_GREY, align=PP_ALIGN.CENTER)

# Per-class table
rect(s, 0.35, 3.15, 12.6, 2.85, DARK_CARD)
col_hdrs = ["Class", "Precision", "Recall", "F1-Score", "Support", "Key Insight"]
c_x = [0.45, 2.55, 3.65, 4.75, 5.85, 6.7]
c_w = [2.0,  1.0,  1.0,  1.0,  1.0,  6.2]

rect(s, 0.35, 3.15, 12.6, 0.38, SKY_BLUE)
for j, h in enumerate(col_hdrs):
    txbox(s, c_x[j], 3.18, c_w[j], 0.32, h, size=11, bold=True, color=WHITE)

class_rows = [
    (SKY_BLUE,  "Low_to_Medium_Rain",  "54.0%", "72.0%", "61.7%", "150",
     "High recall — stratiform cloud decks reliably detected"),
    (ROSE,      "Medium_to_Heavy_Rain","69.2%", "38.7%", "49.7%", "93",
     "High precision — when predicted, it is usually correct"),
    (GREEN,     "No_to_Low_Rain",      "60.5%", "56.5%", "58.4%", "138",
     "Clear sky well-separated from severe convection"),
]
for i, (c, cls, pr, re, f1, sup, insight) in enumerate(class_rows):
    ry = 3.57 + i * 0.65
    bg_c = DARK_CARD if i % 2 == 0 else RGBColor(0x12,0x1F,0x40)
    rect(s, 0.35, ry, 12.6, 0.62, bg_c)
    rect(s, 0.35, ry, 0.06, 0.62, c)
    for j, val in enumerate([cls, pr, re, f1, sup, insight]):
        is_bold = j == 0
        col = c if j == 0 else (WHITE if j < 5 else LIGHT_GREY)
        txbox(s, c_x[j], ry+0.1, c_w[j], 0.42,
              val, size=11 if j != 5 else 10, bold=is_bold, color=col)

txbox(s, 0.35, 6.55, 12.6, 0.35,
      "★  Baseline random-guess accuracy = 33.3%  ·  Model achieves 58.27% — a genuine physical signal extracted from optical cloud textures",
      size=10, italic=True, color=LIGHT_GREY)

slide_number(s, 7)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 8 — Confusion Matrix
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "CONFUSION MATRIX")
slide_title(s, "Confusion Matrix — Test Set (N = 381)",
            "Rows = Actual Class  ·  Columns = Predicted Class")
hline(s, 1.55)

# Confusion matrix grid
cm_labels = ["Low/Med\nRain", "Med/Heavy\nRain", "No/Low\nRain"]
cm_data = [
    [108, 10, 32],
    [38,  36, 19],
    [54,   6, 78],
]
cm_colors = [
    [GREEN, AMBER, ROSE],
    [ROSE,  GREEN, AMBER],
    [ROSE,  AMBER, GREEN],
]
cell_size = 1.55
start_cx = 1.5
start_cy = 1.85

# Col headers
txbox(s, start_cx, 1.7, cell_size*3, 0.3, "PREDICTED →",
      size=11, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
for j, lbl in enumerate(cm_labels):
    rect(s, start_cx + j*cell_size, 2.0, cell_size, 0.55, SKY_BLUE)
    txbox(s, start_cx + j*cell_size+0.05, 2.02,
          cell_size-0.1, 0.5, lbl, size=10, bold=True, color=WHITE,
          align=PP_ALIGN.CENTER)

# Row headers + cells
for i, (row, lbl) in enumerate(zip(cm_data, cm_labels)):
    ry = 2.58 + i * cell_size
    rect(s, 0.4, ry, 1.05, cell_size, SKY_BLUE)
    txbox(s, 0.42, ry+0.3, 1.0, 0.95, lbl,
          size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for j, (val, c) in enumerate(zip(row, cm_colors[i])):
        cx = start_cx + j * cell_size
        is_diag = (i == j)
        fill_c = RGBColor(0x00, 0x70, 0x50) if is_diag else DARK_CARD
        rect(s, cx, ry, cell_size, cell_size, fill_c)
        if is_diag:
            rect(s, cx, ry, cell_size, 0.06, GREEN)
        txbox(s, cx, ry+0.3, cell_size, 0.7, str(val),
              size=30, bold=True, color=WHITE if not is_diag else CYAN,
              align=PP_ALIGN.CENTER)
        pct = f"{val/381*100:.1f}%"
        txbox(s, cx, ry+0.95, cell_size, 0.35, pct,
              size=10, color=LIGHT_GREY, align=PP_ALIGN.CENTER)

# ACTUAL label
txbox(s, 0.05, 3.5, 0.5, 2.5, "A\nC\nT\nU\nA\nL\n↓",
      size=10, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

# Insights
insights = [
    (GREEN, "Strong Stratiform Detection",
     "108/150 Low/Med Rain correctly identified\nRecall = 72.0%"),
    (AMBER, "Convective Precision",
     "When Med/Heavy predicted → 69.2% correct\nLow false heavy-rain alarm rate"),
    (SKY_BLUE, "Clear Sky Separation",
     "Only 6 dry-sky specimens misclassified\nas severe convective (1.6%)"),
]
for i, (c, hdr, body) in enumerate(insights):
    iy = 1.9 + i * 1.65
    rect(s, 6.5, iy, 6.4, 1.45, DARK_CARD)
    rect(s, 6.5, iy, 6.4, 0.06, c)
    txbox(s, 6.6, iy+0.1, 6.2, 0.35, hdr, size=13, bold=True, color=c)
    txbox(s, 6.6, iy+0.5, 6.2, 0.85, body, size=12, color=WHITE)

slide_number(s, 8)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 9 — System Architecture
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "SYSTEM ARCHITECTURE")
slide_title(s, "SKYsense AI — Full System Architecture",
            "5-tier design: Client → Security → Django → AI Engine → Database")
hline(s, 1.55)

layers = [
    (CYAN,     "CLIENT LAYER",      0.35, 1.75, 12.6, 0.9,
     "Web Browser (UI)    ·    Simulated IoT Client (CLI)    ·    Future ESP32 / Raspberry Pi Camera Node"),
    (AMBER,    "SECURITY LAYER",    0.35, 2.73, 12.6, 0.9,
     "CSRF Middleware   ·   File Size Limit (15 MB)   ·   MIME Whitelist (JPEG/PNG)   ·   Pillow Binary Verification   ·   UUID Storage"),
    (SKY_BLUE, "DJANGO APP LAYER",  0.35, 3.71, 12.6, 0.9,
     "Authentication Views   ·   Prediction Views   ·   Demo Engine (/demo/)   ·   REST API (/api/predict/)   ·   Dashboard & History"),
    (GREEN,    "AI INFERENCE ENGINE", 0.35, 4.69, 12.6, 0.9,
     "RainfallInferenceService (Singleton)   ·   Bilinear Resize 256×256   ·   Xception CNN   ·   Softmax → 3-Class Vector"),
    (ROSE,     "PERSISTENCE LAYER", 0.35, 5.67, 12.6, 0.9,
     "SQLite3 / PostgreSQL   ·   PredictionRecord ORM   ·   UUID Media Storage   ·   Evaluation Metrics JSON"),
]
for c, name, lx, ly, lw, lh, content in layers:
    rect(s, lx, ly, lw, lh, DARK_CARD)
    rect(s, lx, ly, lw, 0.05, c)
    rect(s, lx, ly, 2.2, lh, RGBColor(0x0D,0x15,0x30))
    txbox(s, lx+0.1, ly+0.22, 2.0, 0.5, name,
          size=11, bold=True, color=c, align=PP_ALIGN.CENTER)
    txbox(s, lx+2.35, ly+0.2, 10.1, 0.55, content,
          size=12, color=WHITE)
    if ly < 5.5:
        txbox(s, 6.3, ly+lh, 0.7, 0.12, "▼",
              size=10, color=CYAN, align=PP_ALIGN.CENTER)

slide_number(s, 9)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 10 — Web Application Features
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "WEB APPLICATION")
slide_title(s, "Django Web Portal — Feature Overview",
            "Authenticated · Responsive · Chart.js Analytics · Defense-in-Depth Security")
hline(s, 1.55)

features = [
    (SKY_BLUE, "🏠", "/", "Landing Portal",
     "Technology overview, pipeline diagram, live metric badges"),
    (GREEN, "🔐", "/register/ & /login/", "Authentication",
     "Secure register, login, logout, profile management"),
    (CYAN, "🌤️", "/predict/", "Cloud Upload & AI Analysis",
     "Drag-and-drop JPEG/PNG upload → real Xception inference"),
    (AMBER, "🎯", "/demo/", "Demo Mode",
     "Projector-optimized, 12 pre-loaded specimens, laser scan animation"),
    (SKY_BLUE, "📊", "/dashboard/", "Analytics Dashboard",
     "Chart.js: daily volume, class distribution, confidence histogram"),
    (GREEN, "📜", "/history/", "Prediction History",
     "Searchable, filterable audit log with UUID tracking"),
    (ROSE, "📈", "/performance/", "Model Performance",
     "Live test accuracy (58.27%), loss, F1, confusion matrix"),
    (AMBER, "🔌", "/api/predict/", "REST API",
     "POST multipart/form-data → JSON response, IoT-ready"),
]
for i, (c, icon, url, name, desc) in enumerate(features):
    row = i // 4
    col = i % 4
    bx = 0.35 + col * 3.25
    by = 1.75 + row * 2.7
    rect(s, bx, by, 3.1, 2.45, DARK_CARD)
    rect(s, bx, by, 3.1, 0.06, c)
    txbox(s, bx+0.1, by+0.1, 3.0, 0.5, f"{icon}  {name}",
          size=13, bold=True, color=c)
    txbox(s, bx+0.1, by+0.55, 3.0, 0.3, url,
          size=10, color=CYAN, italic=True)
    txbox(s, bx+0.1, by+0.9, 2.9, 1.4, desc,
          size=11, color=WHITE)

slide_number(s, 10)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 11 — REST API
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "REST API")
slide_title(s, "REST API — POST /api/predict/",
            "Stateless · multipart/form-data · JSON Response · IoT-Forward Compatible")
hline(s, 1.55)

# Request side
rect(s, 0.35, 1.7, 6.0, 5.5, DARK_CARD)
txbox(s, 0.45, 1.75, 5.8, 0.4, "REQUEST", size=13, bold=True, color=SKY_BLUE)
req_lines = [
    "POST /api/predict/ HTTP/1.1",
    "Content-Type: multipart/form-data",
    "",
    "image         ← REQUIRED (JPEG/PNG, ≤ 15 MB)",
    "device_id     ← optional: 'RPI-NODE-01'",
    "latitude      ← optional: float (-90 to +90)",
    "longitude     ← optional: float (-180 to +180)",
    "temperature   ← optional: °C float",
    "humidity      ← optional: % float (0–100)",
]
for i, l in enumerate(req_lines):
    color = CYAN if i < 2 else (SKY_BLUE if "REQUIRED" in l else WHITE)
    bold  = i < 2
    txbox(s, 0.45, 2.2+i*0.48, 5.8, 0.45, l,
          size=11, bold=bold, color=color)

# Response side
rect(s, 6.65, 1.7, 6.3, 5.5, DARK_CARD)
txbox(s, 6.75, 1.75, 6.1, 0.4, "JSON RESPONSE", size=13, bold=True, color=GREEN)
resp_lines = [
    '{',
    '  "success": true,',
    '  "prediction": "Medium_to_Heavy_Rain",',
    '  "confidence": 0.503,',
    '  "probabilities": {',
    '    "Low_to_Medium_Rain": 0.295,',
    '    "Medium_to_Heavy_Rain": 0.503,',
    '    "No_to_Low_Rain": 0.202',
    '  },',
    '  "timestamp": "2026-09-24T06:00:00Z",',
    '  "model_version": "Xception-v1.0",',
    '  "processing_time_ms": 38.4',
    '}',
]
for i, l in enumerate(resp_lines):
    color = GREEN if '"success"' in l or '"prediction"' in l else (CYAN if '"confidence"' in l else WHITE)
    txbox(s, 6.75, 2.2+i*0.38, 6.1, 0.36, l,
          size=10, color=color)

slide_number(s, 11)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 12 — Security
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "SECURITY")
slide_title(s, "Defense-in-Depth Security Architecture",
            "OWASP-aligned hardening · manage.py check --deploy: 0 issues")
hline(s, 1.55)

controls = [
    (GREEN,    "File Upload Validation",
     "15 MB hard limit · JPEG/PNG whitelist · Pillow binary header verification"),
    (SKY_BLUE, "UUID-Based File Storage",
     "Uploads stored as random UUID hex filenames — prevents path injection and overwrites"),
    (CYAN,     "CSRF Protection",
     "CsrfViewMiddleware active · CSRF_COOKIE_HTTPONLY=True · SameSite=Lax"),
    (AMBER,    "Path Traversal Prevention",
     ".resolve().relative_to() bounds checking on all sample image access"),
    (GREEN,    "Secure Cookies (Production)",
     "SESSION_COOKIE_SECURE=True · CSRF_COOKIE_SECURE=True when DJANGO_ENV=production"),
    (SKY_BLUE, "HSTS (1 Year)",
     "SECURE_HSTS_SECONDS=31536000 · includeSubDomains · preload"),
    (ROSE,     "No Stack Trace Leaks",
     "Custom 400/403/404/500 error views · DEBUG=False in production"),
    (AMBER,    "Secret Management",
     "DJANGO_SECRET_KEY from .env · .env excluded by .gitignore · .env.example has only placeholders"),
    (GREEN,    "Authorization Isolation",
     "Users only see own predictions · Admin inspection privilege for staff"),
    (CYAN,     "WhiteNoise Static Serving",
     "CompressedManifestStaticFilesStorage · Immutable cache headers · No separate CDN required"),
]
for i, (c, hdr, body) in enumerate(controls):
    row, col = i // 2, i % 2
    bx = 0.35 + col * 6.5
    by = 1.75 + row * 1.12
    rect(s, bx, by, 6.25, 1.0, DARK_CARD)
    rect(s, bx, by, 6.25, 0.06, c)
    txbox(s, bx+0.1, by+0.1, 6.1, 0.35, hdr, size=12, bold=True, color=c)
    txbox(s, bx+0.1, by+0.45, 6.1, 0.48, body, size=11, color=WHITE)

slide_number(s, 12)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 13 — IoT Architecture (Future Work)
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "IoT INTEGRATION")
slide_title(s, "Future IoT Integration Architecture",
            "Software interface complete · Physical hardware is explicitly future research work")
hline(s, 1.55)

# IMPORTANT disclaimer box
rect(s, 0.35, 1.72, 12.6, 0.55, RGBColor(0x4A, 0x20, 0x00))
rect(s, 0.35, 1.72, 12.6, 0.06, AMBER)
txbox(s, 0.45, 1.75, 12.4, 0.45,
      "⚠  ACADEMIC INTEGRITY NOTICE: Physical IoT hardware was NOT built as part of this thesis. "
      "The software architecture is complete and validated using a simulated edge client.",
      size=11, bold=True, color=AMBER)

# Architecture flow
nodes = [
    (DEEP_BLUE, SKY_BLUE, "EDGE CAMERA\n(Future)\nESP32-CAM\nRaspberry Pi\nJetson Nano",  1.0, 2.45),
    (DEEP_BLUE, GREEN,    "IoT SIMULATOR\n(Current)\nSimulate Device\nPython CLI\nCLI Validated",     4.15, 2.45),
    (DEEP_BLUE, CYAN,     "REST API\n(Complete)\nPOST /api/predict/\nmultipart/form-data\nJSON Response",  7.3, 2.45),
    (DEEP_BLUE, AMBER,    "AI ENGINE\n(Complete)\nXception CNN\nSingleton Service\n38ms Inference",    10.45, 2.45),
]
for fill, border, text, nx, ny in nodes:
    rect(s, nx, ny, 2.8, 2.6, fill)
    rect(s, nx, ny, 2.8, 0.07, border)
    txbox(s, nx+0.1, ny+0.15, 2.6, 2.3, text,
          size=12, bold=False, color=WHITE, align=PP_ALIGN.CENTER)
    if nx < 10.0:
        txbox(s, nx+2.82, ny+1.0, 0.3, 0.6, "→",
              size=18, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

# Telemetry fields
rect(s, 0.35, 5.2, 12.6, 1.95, DARK_CARD)
txbox(s, 0.45, 5.25, 12.4, 0.4,
      "Telemetry Fields Supported by /api/predict/ Today",
      size=13, bold=True, color=CYAN)
fields = [
    ("device_id", "Edge node identifier", "RPI-NODE-01"),
    ("latitude",  "GPS Latitude",         "28.6139°N"),
    ("longitude", "GPS Longitude",        "77.2090°E"),
    ("temperature","Ambient temperature", "28.4°C"),
    ("humidity",  "Relative humidity",    "81.2%"),
]
for i, (f, desc, ex) in enumerate(fields):
    cx = 0.45 + i * 2.5
    txbox(s, cx, 5.68, 2.4, 0.3, f, size=11, bold=True, color=SKY_BLUE)
    txbox(s, cx, 5.98, 2.4, 0.25, desc, size=10, color=WHITE)
    txbox(s, cx, 6.23, 2.4, 0.3, f"e.g. {ex}", size=10, color=LIGHT_GREY, italic=True)

txbox(s, 0.35, 7.1, 12.6, 0.3,
      "All telemetry fields are optional and do not affect current web-upload or demo functionality.",
      size=10, italic=True, color=LIGHT_GREY)

slide_number(s, 13)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 14 — IoT Simulator Demo
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "IoT SIMULATOR DEMO")
slide_title(s, "Simulated IoT Edge Client — Live Demonstration",
            "Proves forward-compatibility of the REST API without requiring physical hardware")
hline(s, 1.55)

# Command box
rect(s, 0.35, 1.75, 12.6, 0.6, RGBColor(0x05, 0x10, 0x20))
rect(s, 0.35, 1.75, 12.6, 0.06, GREEN)
txbox(s, 0.45, 1.8, 12.4, 0.48,
      "python iot_simulator/simulate_device.py demo_images/medium_heavy/Cb_Cb-N002.jpg",
      size=14, bold=True, color=GREEN)

# Output terminal
rect(s, 0.35, 2.5, 7.5, 4.65, RGBColor(0x05, 0x0A, 0x18))
rect(s, 0.35, 2.5, 7.5, 0.06, SKY_BLUE)
txbox(s, 0.45, 2.55, 7.3, 0.3, "Terminal Output", size=10, bold=True, color=SKY_BLUE)
terminal_lines = [
    ("════════════════════════════════════", WHITE),
    ("   SIMULATED IoT DEVICE CLIENT", CYAN),
    ("════════════════════════════════════", WHITE),
    ("[!] NOTICE: SIMULATED IoT DEVICE", AMBER),
    ("[INFO] Device ID:    SIM-ESP32-CAM-01", WHITE),
    ("[INFO] Station:      Mountain Weather Node Alpha", WHITE),
    ("[INFO] Coordinates:  Lat 28.6139, Lon 77.2090", WHITE),
    ("[INFO] Sensors:      Temp 28.4°C, Humidity 81.2%", WHITE),
    ("[INFO] POST → http://127.0.0.1:8000/api/predict/", SKY_BLUE),
    ("[SUCCESS] HTTP 200 OK in 42.1 ms", GREEN),
    ("────────────────────────────────────", WHITE),
    ("Verdict:  Medium_to_Heavy_Rain", ROSE),
    ("Confidence:  50.3%", CYAN),
    ("Process time:  42.1 ms", WHITE),
]
for i, (line, c) in enumerate(terminal_lines):
    txbox(s, 0.45, 2.9+i*0.29, 7.3, 0.28, line, size=10, color=c)

# Explanation
explanations = [
    (SKY_BLUE, "What It Demonstrates",
     "The REST API (/api/predict/) accepts cloud\nimages from any HTTP client — browser,\nPython script, or future microcontroller."),
    (GREEN,    "What Is Simulated",
     "Device metadata (GPS, temperature, humidity)\nare randomly generated to simulate sensor\nreadings from a real field station."),
    (AMBER,    "What Is Real",
     "The HTTP POST, the AI inference, the\njson response, and the database record\nare all real — not mocked."),
    (ROSE,     "Future Hardware",
     "When ESP32-CAM or RPi nodes are deployed,\nthey use this exact same API call with\nno server-side changes required."),
]
for i, (c, hdr, body) in enumerate(explanations):
    by = 2.5 + i * 1.13
    rect(s, 8.1, by, 5.15, 1.05, DARK_CARD)
    rect(s, 8.1, by, 5.15, 0.06, c)
    txbox(s, 8.2, by+0.1, 5.0, 0.3, hdr, size=12, bold=True, color=c)
    txbox(s, 8.2, by+0.45, 5.0, 0.58, body, size=11, color=WHITE)

slide_number(s, 14)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 15 — Limitations & Scientific Honesty
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "SCIENTIFIC LIMITATIONS")
slide_title(s, "Acknowledged Limitations & Scientific Boundaries",
            "Transparent reporting is fundamental to academic integrity")
hline(s, 1.55)

limitations = [
    (ROSE, "2D Optical Base Observation",
     "Ground cameras capture only the cloud base. They cannot measure "
     "vertical cloud thickness, cloud-top height, or Liquid Water Path (LWP) — "
     "critical parameters for quantifying rainfall rate."),
    (AMBER, "No Temporal Dynamics",
     "A single-frame snapshot cannot reveal cloud velocity, growth trajectory, "
     "or barometric convergence. A developing Cb and a decaying Cb may be visually identical."),
    (SKY_BLUE, "Virga Phenomenon",
     "Optical models classify based on droplet morphology. In arid/high-temperature "
     "regions, precipitation may evaporate in dry sub-cloud air before reaching the ground."),
    (CYAN, "Illumination Sensitivity",
     "Low solar angles at dawn/dusk and direct glare can skew RGB channel intensity "
     "distributions, reducing classification confidence."),
    (ROSE, "Nighttime Inoperability",
     "Standard RGB sensors cannot function at night. Infrared (IR) or thermal "
     "imaging would be required for 24-hour coverage."),
    (AMBER, "Dataset Geographic Bias",
     "CCSN images are sourced from limited geographic regions. Performance may "
     "degrade on cloud formations from tropical, polar, or desert climates."),
]
for i, (c, hdr, body) in enumerate(limitations):
    row, col = i // 2, i % 2
    bx = 0.35 + col * 6.5
    by = 1.75 + row * 1.75
    rect(s, bx, by, 6.25, 1.6, DARK_CARD)
    rect(s, bx, by, 6.25, 0.06, c)
    txbox(s, bx+0.1, by+0.12, 6.1, 0.38, hdr, size=13, bold=True, color=c)
    txbox(s, bx+0.1, by+0.54, 6.1, 0.95, body, size=11, color=WHITE)

txbox(s, 0.35, 7.1, 12.6, 0.3,
      "★  58.27% accuracy on a 3-class problem (33.3% random baseline) represents genuine physical signal. Full 95%+ precision requires multi-sensor fusion.",
      size=10, italic=True, color=LIGHT_GREY)

slide_number(s, 15)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 16 — Future Development Roadmap
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "FUTURE ROADMAP")
slide_title(s, "Future Development Roadmap",
            "Three research phases planned beyond the current software thesis")
hline(s, 1.55)

phases = [
    (SKY_BLUE, "PHASE 2", "2026–2027",
     "IoT Hardware Deployment",
     [
         "Manufacture ESP32-CAM + BME280 (temp/humidity) field units",
         "Solar-powered weatherproof enclosures for remote sites",
         "Automated cellular/LoRa telemetry to central Django server",
         "Multi-node geographic coverage mapping",
         "Real-time rainfall correlation with ground truth gauges",
     ]),
    (GREEN, "PHASE 3", "2027–2028",
     "Multi-Sensor Fusion & Model Improvement",
     [
         "Add atmospheric pressure (BMP388) and wind sensors",
         "Temporal sequence modeling (LSTM / ConvLSTM on video frames)",
         "Multi-modal fusion: optical + meteorological time series",
         "Larger, geographically diverse dataset collection",
         "Target accuracy improvement to 75%+ with fusion features",
     ]),
    (AMBER, "PHASE 4", "2028+",
     "Production Network & Research Publication",
     [
         "Deploy 20+ sensor nodes across target region",
         "Publish peer-reviewed results with ground-truth validation",
         "Open-source dataset contribution to meteorological community",
         "Integration with civil authority early warning systems",
         "Mobile application for real-time public rainfall alerts",
     ]),
]
for i, (c, phase, yr, title, pts) in enumerate(phases):
    bx = 0.35 + i * 4.35
    rect(s, bx, 1.75, 4.15, 5.35, DARK_CARD)
    rect(s, bx, 1.75, 4.15, 0.07, c)
    txbox(s, bx+0.12, 1.8, 4.0, 0.38, f"{phase}  ({yr})",
          size=11, bold=True, color=c)
    txbox(s, bx+0.12, 2.18, 4.0, 0.5, title,
          size=15, bold=True, color=WHITE)
    hline(s, 2.73, c, l=bx+0.12, w=3.85, thickness=0.025)
    for j, pt in enumerate(pts):
        txbox(s, bx+0.15, 2.82 + j*0.55, 3.9, 0.5,
              f"▸  {pt}", size=11, color=WHITE)

# Current status banner
rect(s, 0.35, 7.05, 12.6, 0.38, RGBColor(0x00, 0x40, 0x28))
rect(s, 0.35, 7.05, 12.6, 0.05, GREEN)
txbox(s, 0.45, 7.1, 12.4, 0.3,
      "✅  CURRENT STATUS (Phase 1 Complete):  Software platform · Trained Xception model · Validated REST API · "
      "IoT Simulator · 87 Automated Tests · Production Deployment Documentation",
      size=10, bold=True, color=GREEN)

slide_number(s, 16)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 17 — Technology Stack Summary
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
section_badge(s, "TECHNOLOGY STACK")
slide_title(s, "Verified Technology Stack",
            "Every dependency verified against installed venv · python-pptx generated this slide")
hline(s, 1.55)

stack_groups = [
    (SKY_BLUE, "Deep Learning", [
        ("TensorFlow", "2.21.0"),
        ("Keras",      "3.15.1"),
        ("NumPy",      "2.4.6"),
        ("h5py",       "3.14.0"),
    ]),
    (GREEN, "Web Framework", [
        ("Django",     "5.2.17"),
        ("WhiteNoise", "6.12.0"),
        ("asgiref",    "3.12.1"),
        ("sqlparse",   "0.6.0"),
    ]),
    (CYAN, "Computer Vision", [
        ("Pillow",     "12.3.0"),
        ("scikit-learn","1.9.1"),
        ("scipy",      "1.17.1"),
        ("matplotlib", "3.11.2"),
    ]),
    (AMBER, "Data & Networking", [
        ("pandas",     "3.0.6"),
        ("seaborn",    "0.13.2"),
        ("requests",   "2.34.2"),
        ("kagglehub",  "1.0.2"),
    ]),
]
for i, (c, group_name, pkgs) in enumerate(stack_groups):
    bx = 0.35 + i * 3.25
    rect(s, bx, 1.75, 3.1, 4.0, DARK_CARD)
    rect(s, bx, 1.75, 3.1, 0.07, c)
    txbox(s, bx+0.1, 1.8, 3.0, 0.4, group_name,
          size=13, bold=True, color=c)
    for j, (pkg, ver) in enumerate(pkgs):
        ry = 2.28 + j * 0.72
        txbox(s, bx+0.12, ry, 1.6, 0.35, pkg,
              size=12, bold=True, color=WHITE)
        txbox(s, bx+1.7, ry, 1.3, 0.35, f"v{ver}",
              size=11, color=LIGHT_GREY)

# Non-Python stack
rect(s, 0.35, 5.95, 12.6, 1.25, DARK_CARD)
txbox(s, 0.45, 6.0, 12.4, 0.4,
      "Frontend, Infrastructure & Tools", size=13, bold=True, color=CYAN)
others = [
    ("HTML5 + CSS3 + Bootstrap 5", "Responsive UI"),
    ("Chart.js 4.x", "Interactive data visualization"),
    ("SQLite 3 / PostgreSQL", "Relational database"),
    ("Nginx + Gunicorn", "Production server stack"),
    ("Let's Encrypt (Certbot)", "TLS/SSL certificates"),
    ("Python-pptx", "This presentation"),
]
for i, (tech, desc) in enumerate(others):
    cx = 0.45 + i * 2.15
    txbox(s, cx, 6.42, 2.1, 0.35, tech, size=11, bold=True, color=SKY_BLUE)
    txbox(s, cx, 6.77, 2.1, 0.3, desc, size=10, color=WHITE)

slide_number(s, 17)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 18 — Conclusion
# ─────────────────────────────────────────────────────────────────────────────
s = new_slide(); bg(s)
rect(s, 0, 0, 13.33, 2.0, DEEP_BLUE)
rect(s, 0, 1.97, 13.33, 0.06, CYAN)

txbox(s, 0.5, 0.25, 12.3, 0.4, "CONCLUSION",
      size=13, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
txbox(s, 0.5, 0.62, 12.3, 1.0, "SKYsense AI",
      size=50, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txbox(s, 0.5, 1.6, 12.3, 0.4,
      "A complete, validated research platform for optical cloud-to-rainfall classification",
      size=14, color=LIGHT_GREY, align=PP_ALIGN.CENTER, italic=True)

# Summary achievement columns
achievements = [
    (GREEN,    "✅", "AI Model",
     "Xception CNN trained\n2-phase transfer learning\n20.8M parameters\nTFLite: 20.5 MB"),
    (SKY_BLUE, "✅", "Web Platform",
     "Django 5.2\n87 tests passing\nDemo mode\nProduction-ready"),
    (CYAN,     "✅", "REST API",
     "POST /api/predict/\nIoT-compatible\nJSON response\nFuture-ready"),
    (AMBER,    "📊", "Verified Results",
     "58.27% test accuracy\n0.9103 loss\n56.60% macro F1\n381 specimens"),
    (ROSE,     "🔭", "Future Work",
     "ESP32 hardware\nMulti-sensor fusion\nGeographic expansion\nPublication"),
]
for i, (c, icon, hdr, body) in enumerate(achievements):
    bx = 0.35 + i * 2.55
    rect(s, bx, 2.2, 2.4, 2.9, DARK_CARD)
    rect(s, bx, 2.2, 2.4, 0.06, c)
    txbox(s, bx, 2.28, 2.4, 0.5, f"{icon}  {hdr}",
          size=13, bold=True, color=c, align=PP_ALIGN.CENTER)
    txbox(s, bx+0.1, 2.82, 2.2, 2.1, body,
          size=12, color=WHITE, align=PP_ALIGN.CENTER)

# Key message
rect(s, 0.35, 5.3, 12.6, 1.1, RGBColor(0x00, 0x2A, 0x4A))
rect(s, 0.35, 5.3, 12.6, 0.06, CYAN)
txbox(s, 0.45, 5.38, 12.4, 0.45,
      "\"SKYsense AI demonstrates that affordable edge-deployable deep learning can extract genuine precipitation signals "
      "from standard ground-level cloud photography — achieving 58.27% accuracy versus a 33.3% random baseline.\"",
      size=13, italic=True, color=WHITE)
txbox(s, 0.45, 5.88, 12.4, 0.35,
      "The software architecture, trained model, validated API, and IoT-ready interface collectively form a complete research prototype.",
      size=12, color=LIGHT_GREY)

hline(s, 6.6, CYAN)
txbox(s, 0.5, 6.65, 8.0, 0.35,
      "Thank you for your attention. Questions welcome.",
      size=16, bold=True, color=WHITE)
txbox(s, 8.5, 6.65, 4.3, 0.35,
      "github.com/Joseph101Gitup/SkySense-AI",
      size=12, color=CYAN, align=PP_ALIGN.RIGHT, italic=True)

slide_number(s, 18)


# ─────────────────────────────────────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────────────────────────────────────
prs.save(str(OUTPUT_PATH))
print(f"\n[OK] Presentation saved to:\n    {OUTPUT_PATH}")
print(f"     Slides: 18  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
